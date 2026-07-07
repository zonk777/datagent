from __future__ import annotations

import base64
import html
import json
import math
import os
import re
import tempfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Image as PdfImage
from reportlab.platypus import LongTable, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from ..db import connect


FONT_CANDIDATES = [
    Path("C:/Windows/Fonts/Deng.ttf"),
    Path("C:/Windows/Fonts/msyh.ttc"),
    Path("C:/Windows/Fonts/simhei.ttf"),
    Path("C:/Windows/Fonts/simsun.ttc"),
]


@dataclass
class ReportData:
    session: dict[str, Any]
    assistant_message: dict[str, Any]
    question: str
    payload: dict[str, Any]


def load_report_data(session_id: str) -> ReportData:
    with connect() as conn:
        session = conn.execute("SELECT * FROM sessions WHERE id = %s", (session_id,)).fetchone()
        if not session:
            raise LookupError("会话不存在")
        assistant = conn.execute(
            """SELECT id, content, payload, created_at
               FROM messages
               WHERE session_id = %s AND role = 'assistant' AND payload IS NOT NULL
               ORDER BY id DESC LIMIT 1""",
            (session_id,),
        ).fetchone()
        if not assistant:
            raise LookupError("该会话还没有可导出的分析结果")
        user = conn.execute(
            """SELECT content
               FROM messages
               WHERE session_id = %s AND role = 'user' AND id < %s
               ORDER BY id DESC LIMIT 1""",
            (session_id, assistant["id"]),
        ).fetchone()
    payload = json.loads(assistant["payload"])
    return ReportData(
        session=dict(session),
        assistant_message=dict(assistant),
        question=user["content"] if user else payload.get("effective_question", ""),
        payload=payload,
    )


def _safe_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value)


def _h(value: Any) -> str:
    return html.escape(_safe_text(value))


def _report_title(data: ReportData) -> str:
    return _safe_text(data.payload.get("chart", {}).get("title") or data.session.get("title") or "数据智能体分析报告")


ALLOWED_CHART_TYPES = {"bar", "line", "pie", "scatter", "area", "radar", "none"}
CN_NUMBERS = ["零", "一", "二", "三", "四", "五", "六", "七", "八", "九", "十"]


def _section_heading(index: int, title: str) -> str:
    prefix = CN_NUMBERS[index] if index < len(CN_NUMBERS) else str(index)
    return f"{prefix}、{title}"


def apply_chart_options(data: ReportData, raw_options: str | dict[str, Any] | None) -> ReportData:
    """Apply per-section chart selections supplied by the frontend export URL."""
    if not raw_options:
        return data
    try:
        options = json.loads(raw_options) if isinstance(raw_options, str) else raw_options
    except (TypeError, json.JSONDecodeError):
        return data
    if not isinstance(options, dict):
        return data

    choices = options.get("sections") or []
    if not isinstance(choices, list):
        return data

    by_index: dict[int, str] = {}
    by_id: dict[str, str] = {}
    for choice in choices:
        if not isinstance(choice, dict):
            continue
        chart_type = _safe_text(choice.get("type")).strip()
        if chart_type not in ALLOWED_CHART_TYPES:
            continue
        try:
            by_index[int(choice.get("index"))] = chart_type
        except (TypeError, ValueError):
            pass
        if choice.get("id"):
            by_id[_safe_text(choice.get("id"))] = chart_type

    sections = data.payload.get("chart_sections") or []
    if sections:
        for index, section in enumerate(sections):
            if not isinstance(section, dict):
                continue
            section_id = _safe_text(section.get("id") or f"section-{index}")
            chart_type = by_id.get(section_id) or by_index.get(index)
            if chart_type:
                section.setdefault("chart", {})["type"] = chart_type
        first_type = by_id.get("primary") or by_index.get(0)
        if first_type and data.payload.get("chart"):
            data.payload["chart"]["type"] = first_type
    else:
        chart_type = by_id.get("primary") or by_index.get(0)
        if chart_type and data.payload.get("chart"):
            data.payload["chart"]["type"] = chart_type
    return data


def _chart_payloads(payload: dict[str, Any]) -> list[dict[str, Any]]:
    sections = payload.get("chart_sections") or []
    output: list[dict[str, Any]] = []
    if isinstance(sections, list) and sections:
        for index, section in enumerate(sections, 1):
            if not isinstance(section, dict):
                continue
            chart = dict(section.get("chart") or {})
            section_title = section.get("title") or chart.get("title") or f"图表 {index}"
            chart.setdefault("title", section_title)
            rows = section.get("rows") or []
            columns = section.get("columns") or []
            if not rows or not columns or chart.get("type") == "none":
                continue
            output.append(
                {
                    "title": section_title,
                    "columns": columns,
                    "rows": rows,
                    "chart": chart,
                    "description": section.get("description") or "",
                    "insights": section.get("insights") if isinstance(section.get("insights"), list) else [],
                }
            )
    elif payload.get("rows") and (payload.get("chart") or {}).get("type") != "none":
        output.append(
            {
                "title": (payload.get("chart") or {}).get("title") or "结果图表",
                "columns": payload.get("columns") or [],
                "rows": payload.get("rows") or [],
                "chart": dict(payload.get("chart") or {}),
                "description": "",
                "insights": [],
            }
        )
    return output


def _is_document_payload(payload: dict[str, Any]) -> bool:
    return _safe_text(payload.get("execution_mode")).startswith(("document", "local-document"))


def _compact_reference_title(title: Any) -> str:
    text = _safe_text(title).strip() or "用户上传文档"
    for sep in ("：", ":"):
        if sep in text:
            text = text.split(sep, 1)[0].strip()
    for suffix in (".pdf", ".docx", ".doc", ".md", ".txt", ".xlsx", ".csv"):
        if text.lower().endswith(suffix):
            text = text[: -len(suffix)].strip()
            break
    return text or "用户上传文档"


def _document_basis_line(payload: dict[str, Any]) -> str:
    refs = payload.get("knowledge_refs") or []
    if refs and isinstance(refs[0], dict):
        return f"分析依据：{_compact_reference_title(refs[0].get('title'))}。"
    return "分析依据：用户上传文档。"


def _top_report_insights(data: ReportData, chart_payloads: list[dict[str, Any]]) -> list[str]:
    payload = data.payload
    if _is_document_payload(payload) and chart_payloads:
        return [_document_basis_line(payload)]
    return [_safe_text(item).strip() for item in payload.get("insights", []) if _safe_text(item).strip()]


def _chart_keywords(chart_payload: dict[str, Any]) -> list[str]:
    chart = chart_payload.get("chart") or {}
    columns = chart_payload.get("columns") or []
    rows = chart_payload.get("rows") or []
    values: list[Any] = [
        chart_payload.get("title"),
        chart_payload.get("description"),
        chart.get("title"),
        chart.get("x_field"),
        chart.get("y_field"),
        chart.get("series_name"),
        chart.get("series_field"),
        *(chart.get("series_fields") or []),
        *columns,
    ]
    for row in rows[:30]:
        if isinstance(row, dict):
            values.extend(row.get(column) for column in columns[:3])
    keywords: list[str] = []
    for value in values:
        text = _safe_text(value).strip()
        for part in re.split(r"[、，,。；;：:\s/()（）\-]+", text):
            part = part.strip()
            if 2 <= len(part) <= 28 and part not in keywords:
                keywords.append(part)
    return keywords


def _chart_insights(chart_payload: dict[str, Any], payload: dict[str, Any]) -> list[str]:
    own = [_safe_text(item).strip() for item in chart_payload.get("insights", []) if _safe_text(item).strip()]
    if own:
        return own
    keywords = _chart_keywords(chart_payload)
    if not keywords:
        return []
    matched = [
        _safe_text(item).strip()
        for item in payload.get("insights", [])
        if _safe_text(item).strip() and any(keyword in _safe_text(item) for keyword in keywords)
    ]
    return matched[:3]


def _overall_insights(payload: dict[str, Any], chart_payloads: list[dict[str, Any]]) -> list[str]:
    insights = [_safe_text(item).strip() for item in payload.get("insights", []) if _safe_text(item).strip()]
    if not insights:
        return []
    assigned: set[str] = set()
    for chart_payload in chart_payloads:
        assigned.update(_chart_insights(chart_payload, payload))
    basis = _document_basis_line(payload) if _is_document_payload(payload) else ""
    remaining = [item for item in insights if item not in assigned and item != basis]
    preferred = [
        item
        for item in remaining
        if any(word in item for word in ("建议", "关注", "风险", "应", "需要", "后续", "优化", "下钻", "复盘", "验证"))
    ]
    selected = preferred or remaining
    if selected:
        return selected[:6]
    if _is_document_payload(payload) and chart_payloads:
        return ["建议结合各图表继续下钻异常指标、增长来源与风险项，并补充业务口径进行交叉验证。"]
    if chart_payloads:
        return ["建议围绕上方图表中的高值、低值、趋势拐点和异常波动继续下钻，结合业务规则验证原因。"]
    return []


def _chart_image_uri(chart_payload: dict[str, Any]) -> str:
    image = build_chart_image(chart_payload)
    if not image:
        return ""
    encoded = base64.b64encode(image.getvalue()).decode("ascii")
    return f"data:image/png;base64,{encoded}"


def build_html_report(data: ReportData) -> str:
    payload = data.payload
    chart_payloads = _chart_payloads(payload)
    insights = "".join(f"<li>{_h(item)}</li>" for item in _top_report_insights(data, chart_payloads))
    references = "".join(
        f"<li><strong>{_h(_compact_reference_title(item.get('title', '')))}</strong> <span>{_h(item.get('category', '知识依据'))}</span></li>"
        for item in payload.get("knowledge_refs", [])
    )
    chart_blocks = ""
    for chart_payload in chart_payloads:
        uri = _chart_image_uri(chart_payload)
        if uri:
            chart_insights = "".join(f"<li>{_h(item)}</li>" for item in _chart_insights(chart_payload, payload))
            insight_block = f"<div class='chart-insights'><h4>该图提示</h4><ul>{chart_insights}</ul></div>" if chart_insights else ""
            chart_blocks += (
                f"<section class='chart-block'><h3>{_h(chart_payload.get('title'))}</h3>"
                f"<img src='{uri}' alt='{_h(chart_payload.get('title'))}' />{insight_block}</section>"
            )
    overall_items = "".join(f"<li>{_h(item)}</li>" for item in _overall_insights(payload, chart_payloads))
    overall_block = f"<h2>总体建议</h2><ul>{overall_items}</ul>" if overall_items else ""
    table = ""
    if payload.get("rows"):
        columns = payload.get("columns", [])
        table_head = "".join(f"<th>{_h(column)}</th>" for column in columns)
        table_rows = "".join(
            "<tr>" + "".join(f"<td>{_h(item.get(column, ''))}</td>" for column in columns) + "</tr>"
            for item in payload["rows"]
        )
        table = f"<h2>查询结果</h2><table><thead><tr>{table_head}</tr></thead><tbody>{table_rows}</tbody></table>"
    sql = f"<h2>执行 SQL</h2><pre>{_h(payload['sql'])}</pre>" if payload.get("sql") else ""
    refs = f"<h2>知识依据</h2><ul>{references}</ul>" if references else ""
    return f"""<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><title>数据智能体分析报告</title>
    <style>
    body{{font-family:Arial,'Microsoft YaHei',sans-serif;max-width:1080px;margin:0 auto;padding:46px 34px;color:#16324f;line-height:1.75;background:linear-gradient(135deg,#f5fbff,#eef8fb)}}
    .report{{background:rgba(255,255,255,.92);border:1px solid #dbe9f2;border-radius:24px;padding:34px 38px;box-shadow:0 24px 60px rgba(38,80,116,.12)}}
    h1{{color:#0b3150;margin:0 0 10px;font-size:30px}} h2{{color:#087ea4;margin-top:28px;border-left:4px solid #19b9c6;padding-left:10px}}
    .meta{{color:#667085;background:#f2f8fc;border-radius:14px;padding:10px 14px}} table{{width:100%;border-collapse:collapse;margin:24px 0;background:white}}
    th,td{{padding:10px 12px;border:1px solid #dbe5ee;text-align:left}} th{{background:#edf8fc;color:#31516a}} pre{{background:#102b40;color:#d9f3f1;border-radius:14px;padding:16px;white-space:pre-wrap;overflow:auto}}
    li{{margin:8px 0}} li span{{color:#718096;font-size:13px;margin-left:8px}}
    .chart-block{{margin:22px 0;padding:18px;border:1px solid #dbe9f2;border-radius:18px;background:#fbfdff}}
    .chart-block h3{{margin:0 0 12px;color:#173a55}} .chart-block img{{width:100%;max-width:980px;border-radius:14px;border:1px solid #e3edf5}}
    .chart-insights{{margin-top:14px;padding:12px 14px;border-radius:14px;background:#f2faf9;border:1px solid #d9eceb}} .chart-insights h4{{margin:0 0 8px;color:#146f82}} .chart-insights ul{{margin:0;padding-left:22px}}
    </style></head>
    <body><main class="report"><h1>{_h(_report_title(data))}</h1>
    <p class="meta">会话编号：{_h(data.session.get("id", ""))} · 类型：{_h(payload.get("intent", ""))} · 生成时间：{_h(data.assistant_message.get("created_at", ""))}</p>
    <p><strong>分析问题：</strong>{_h(data.question)}</p>
    <h2>回答与发现</h2><ul>{insights}</ul><h2>结果图表</h2>{chart_blocks or '<p>本次报告未生成图表。</p>'}{overall_block}{table}{sql}{refs}</main></body></html>"""


def build_markdown_report(data: ReportData) -> bytes:
    payload = data.payload
    chart_payloads = _chart_payloads(payload)
    lines: list[str] = [
        f"# {_report_title(data)}",
        "",
        "## 报告概况",
        "",
        f"- 会话编号：{_safe_text(data.session.get('id', ''))}",
        f"- 分析问题：{_safe_text(data.question)}",
        f"- 分析类型：{_safe_text(payload.get('intent', ''))}",
        f"- 执行模式：{_safe_text(payload.get('execution_mode', ''))}",
        f"- 生成时间：{_safe_text(data.assistant_message.get('created_at', ''))}",
        "",
        "## 回答与关键发现",
        "",
    ]
    for idx, insight in enumerate(_top_report_insights(data, chart_payloads), 1):
        lines.append(f"{idx}. {_safe_text(insight)}")
    lines.extend(["", "## 结果图表", ""])
    if chart_payloads:
        for chart_payload in chart_payloads:
            uri = _chart_image_uri(chart_payload)
            if uri:
                title = _safe_text(chart_payload.get("title") or "图表")
                lines.extend([f"### {title}", "", f"![{title}]({uri})", ""])
                chart_insights = _chart_insights(chart_payload, payload)
                if chart_insights:
                    lines.extend(["**该图提示：**", ""])
                    for idx, insight in enumerate(chart_insights, 1):
                        lines.append(f"{idx}. {_safe_text(insight)}")
                    lines.append("")
    else:
        lines.append("> 用户选择不生成图像，或当前结果没有可视化数据。")
    overall = _overall_insights(payload, chart_payloads)
    if overall:
        lines.extend(["", "## 总体建议", ""])
        for idx, insight in enumerate(overall, 1):
            lines.append(f"{idx}. {_safe_text(insight)}")
    if payload.get("rows"):
        columns = payload.get("columns", [])
        lines.extend(["", "## 查询结果", "", "| " + " | ".join(map(str, columns)) + " |"])
        lines.append("| " + " | ".join("---" for _ in columns) + " |")
        for row in payload.get("rows", [])[:120]:
            lines.append("| " + " | ".join(_safe_text(row.get(column, "")).replace("|", "\\|") for column in columns) + " |")
        if len(payload.get("rows", [])) > 120:
            lines.append(f"\n> 结果共 {len(payload.get('rows', []))} 行，Markdown 仅展示前 120 行。")
    if payload.get("sql"):
        lines.extend(["", "## 执行 SQL", "", "```sql", _safe_text(payload["sql"]), "```"])
    if payload.get("knowledge_refs"):
        lines.extend(["", "## 知识依据", ""])
        for item in payload["knowledge_refs"]:
            lines.append(f"- **{_compact_reference_title(item.get('title', ''))}**（{_safe_text(item.get('category', ''))}）")
    return "\n".join(lines).encode("utf-8")


def _find_font_path() -> Path | None:
    return next((path for path in FONT_CANDIDATES if path.exists()), None)


def _image_font(size: int, bold: bool = False) -> ImageFont.ImageFont:
    path = Path("C:/Windows/Fonts/Dengb.ttf") if bold and Path("C:/Windows/Fonts/Dengb.ttf").exists() else _find_font_path()
    if path:
        return ImageFont.truetype(str(path), size=size)
    return ImageFont.load_default()


def _float_value(value: Any) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _series_key(row: dict[str, Any], series_fields: list[str]) -> str:
    if not series_fields:
        return "结果"
    return " / ".join(_safe_text(row.get(field) or "未分类") for field in series_fields)


def build_chart_image(payload: dict[str, Any]) -> BytesIO | None:
    chart = payload.get("chart") or {}
    rows = payload.get("rows") or []
    chart_type = chart.get("type")
    x_field = chart.get("x_field")
    y_field = chart.get("y_field")
    if not rows or not x_field or not y_field or chart_type == "none":
        return None

    width, height = 1100, 520
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    title_font = _image_font(28, True)
    label_font = _image_font(18)
    small_font = _image_font(15)
    axis_color = (76, 100, 132)
    grid_color = (224, 234, 242)
    colorset = [(22, 119, 255), (16, 183, 196), (91, 117, 231), (245, 158, 11), (16, 163, 127), (239, 71, 111)]
    draw.text((42, 26), _safe_text(chart.get("title") or "分析图表"), font=title_font, fill=(20, 44, 76))
    plot = (78, 95, 1040, 420)

    series_fields = chart.get("series_fields") or ([chart["series_field"]] if chart.get("series_field") else [])
    grouped: dict[str, dict[str, float]] = {}
    x_values: list[str] = []
    for row in rows:
        x = _safe_text(row.get(x_field))
        s = _series_key(row, series_fields)
        if x not in x_values:
            x_values.append(x)
        grouped.setdefault(s, {})[x] = grouped.setdefault(s, {}).get(x, 0) + _float_value(row.get(y_field))
    x_values = x_values[:28]

    if chart_type == "radar":
        radar_labels = x_values[:8]
        values = [value for group in grouped.values() for value in group.values()]
        max_value = max(values) if values else 1
        center_x, center_y, radius = 550, 270, 165
        for ring in range(1, 6):
            pts = []
            current_r = radius * ring / 5
            for idx in range(len(radar_labels)):
                angle = -math.pi / 2 + 2 * math.pi * idx / max(len(radar_labels), 1)
                pts.append((center_x + current_r * math.cos(angle), center_y + current_r * math.sin(angle)))
            if len(pts) > 2:
                draw.polygon(pts, outline=grid_color)
        for idx, label in enumerate(radar_labels):
            angle = -math.pi / 2 + 2 * math.pi * idx / max(len(radar_labels), 1)
            end = (center_x + radius * math.cos(angle), center_y + radius * math.sin(angle))
            draw.line((center_x, center_y, end[0], end[1]), fill=grid_color, width=1)
            draw.text((end[0] - 35, end[1] - 10), label[:8], font=small_font, fill=(93, 111, 132))
        for s_idx, (series, data) in enumerate(list(grouped.items())[:6]):
            pts = []
            for idx, label in enumerate(radar_labels):
                angle = -math.pi / 2 + 2 * math.pi * idx / max(len(radar_labels), 1)
                current_r = radius * (data.get(label, 0) / max(max_value, 1))
                pts.append((center_x + current_r * math.cos(angle), center_y + current_r * math.sin(angle)))
            if len(pts) > 2:
                draw.line(pts + [pts[0]], fill=colorset[s_idx % len(colorset)], width=4)
            lx = 82 + s_idx * 150
            draw.rectangle((lx, 452, lx + 18, 470), fill=colorset[s_idx % len(colorset)])
            draw.text((lx + 26, 448), series[:10], font=small_font, fill=(70, 88, 112))
    elif chart_type == "pie":
        totals = [(x, sum(group.get(x, 0) for group in grouped.values())) for x in x_values[:10]]
        total = sum(value for _, value in totals) or 1
        bbox = (165, 120, 475, 430)
        start = 0.0
        for idx, (label, value) in enumerate(totals):
            extent = value / total * 360
            draw.pieslice(bbox, start, start + extent, fill=colorset[idx % len(colorset)], outline="white", width=2)
            start += extent
            y = 125 + idx * 30
            draw.rectangle((565, y + 3, 585, y + 23), fill=colorset[idx % len(colorset)])
            draw.text((595, y), f"{label}  {value:,.2f}", font=label_font, fill=(40, 58, 82))
    else:
        values = [value for group in grouped.values() for value in group.values()]
        max_value = max(values) if values else 1
        min_value = min(0, min(values) if values else 0)
        span = max(max_value - min_value, 1)
        x1, y1, x2, y2 = plot
        for idx in range(6):
            y = y2 - int((y2 - y1) * idx / 5)
            draw.line((x1, y, x2, y), fill=grid_color, width=1)
            label = min_value + span * idx / 5
            draw.text((10, y - 10), f"{label:,.0f}", font=small_font, fill=(112, 129, 151))
        draw.line((x1, y2, x2, y2), fill=axis_color, width=2)
        draw.line((x1, y1, x1, y2), fill=axis_color, width=2)
        if chart_type in ("line", "area", "scatter"):
            for s_idx, (series, data) in enumerate(list(grouped.items())[:6]):
                pts = []
                for idx, x in enumerate(x_values):
                    px = x1 + int((x2 - x1) * idx / max(len(x_values) - 1, 1))
                    py = y2 - int((data.get(x, 0) - min_value) / span * (y2 - y1))
                    pts.append((px, py))
                if chart_type == "area" and len(pts) > 1:
                    draw.polygon([*pts, (pts[-1][0], y2), (pts[0][0], y2)], fill=(230, 242, 255))
                if chart_type in ("line", "area") and len(pts) > 1:
                    draw.line(pts, fill=colorset[s_idx % len(colorset)], width=4)
                for px, py in pts:
                    dot = 6 if chart_type == "scatter" else 4
                    draw.ellipse((px - dot, py - dot, px + dot, py + dot), fill=colorset[s_idx % len(colorset)])
        else:
            totals = [(x, sum(group.get(x, 0) for group in grouped.values())) for x in x_values]
            bar_w = max(10, int((x2 - x1) / max(len(totals), 1) * 0.62))
            for idx, (x, value) in enumerate(totals):
                cx = x1 + int((x2 - x1) * (idx + 0.5) / max(len(totals), 1))
                top = y2 - int((value - min_value) / span * (y2 - y1))
                draw.rounded_rectangle((cx - bar_w // 2, top, cx + bar_w // 2, y2), radius=5, fill=colorset[idx % len(colorset)])
        label_step = max(1, math.ceil(len(x_values) / 8))
        for idx, x in enumerate(x_values):
            if idx % label_step:
                continue
            px = x1 + int((x2 - x1) * idx / max(len(x_values) - 1, 1))
            draw.text((px - 38, y2 + 12), x[:10], font=small_font, fill=(93, 111, 132))
        for s_idx, series in enumerate(list(grouped.keys())[:6]):
            lx = 82 + s_idx * 150
            draw.rectangle((lx, 452, lx + 18, 470), fill=colorset[s_idx % len(colorset)])
            draw.text((lx + 26, 448), series[:10], font=small_font, fill=(70, 88, 112))

    out = BytesIO()
    img.save(out, format="PNG")
    out.seek(0)
    return out


def _set_run_font(run, size: int | None = None, bold: bool | None = None, color: RGBColor | None = None) -> None:
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = color


def _add_docx_paragraph(doc: Document, text: str, size: int = 10, bold: bool = False) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(_safe_text(text))
    _set_run_font(run, size=size, bold=bold)


def _add_docx_table(doc: Document, headers: list[str], rows: list[list[Any]], max_rows: int = 80) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = _safe_text(header)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                _set_run_font(run, size=9, bold=True, color=RGBColor(11, 37, 69))
    for row in rows[:max_rows]:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = _safe_text(value)
            for paragraph in cells[idx].paragraphs:
                for run in paragraph.runs:
                    _set_run_font(run, size=8)
    if len(rows) > max_rows:
        _add_docx_paragraph(doc, f"注：结果共 {len(rows)} 行，Word 报告仅展示前 {max_rows} 行。", size=9)


def build_docx_report(data: ReportData) -> bytes:
    payload = data.payload
    chart_payloads = _chart_payloads(payload)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("数据智能体分析报告")
    _set_run_font(run, size=22, bold=True, color=RGBColor(11, 37, 69))
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(_report_title(data))
    _set_run_font(run, size=13, bold=True, color=RGBColor(46, 116, 181))

    doc.add_heading("一、报告概况", level=1)
    _add_docx_table(
        doc,
        ["项目", "内容"],
        [
            ["会话编号", data.session["id"]],
            ["分析问题", data.question],
            ["分析类型", payload.get("intent", "")],
            ["执行模式", payload.get("execution_mode", "")],
            ["生成时间", data.assistant_message.get("created_at", "")],
        ],
    )

    doc.add_heading("二、回答与关键发现", level=1)
    for idx, insight in enumerate(_top_report_insights(data, chart_payloads), 1):
        _add_docx_paragraph(doc, f"{idx}. {insight}", size=10)

    if chart_payloads:
        doc.add_heading("三、结果图表", level=1)
        for index, chart_payload in enumerate(chart_payloads, 1):
            _add_docx_paragraph(doc, f"{index}. {_safe_text(chart_payload.get('title') or '图表')}", size=10, bold=True)
            chart_image = build_chart_image(chart_payload)
            if chart_image:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run().add_picture(chart_image, width=Inches(6.4))
            chart_insights = _chart_insights(chart_payload, payload)
            if chart_insights:
                _add_docx_paragraph(doc, "该图提示：", size=10, bold=True)
                for insight_index, insight in enumerate(chart_insights, 1):
                    _add_docx_paragraph(doc, f"{insight_index}. {insight}", size=9)

    next_section = 4 if chart_payloads else 3
    overall = _overall_insights(payload, chart_payloads)
    if overall:
        doc.add_heading(_section_heading(next_section, "总体建议"), level=1)
        next_section += 1
        for idx, insight in enumerate(overall, 1):
            _add_docx_paragraph(doc, f"{idx}. {insight}", size=10)

    if payload.get("rows"):
        doc.add_heading(_section_heading(next_section, "查询结果"), level=1)
        next_section += 1
        columns = payload.get("columns", [])
        _add_docx_table(doc, columns, [[row.get(column, "") for column in columns] for row in payload.get("rows", [])])

    if payload.get("sql"):
        doc.add_heading(_section_heading(next_section, "执行 SQL"), level=1)
        next_section += 1
        _add_docx_paragraph(doc, payload["sql"], size=9)

    if payload.get("knowledge_refs"):
        doc.add_heading(_section_heading(next_section, "知识依据"), level=1)
        _add_docx_table(
            doc,
            ["标题", "类别"],
            [[_compact_reference_title(item.get("title", "")), item.get("category", "")] for item in payload["knowledge_refs"]],
            max_rows=20,
        )

    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _register_pdf_font() -> str:
    font_name = "DataAgentCN"
    if font_name in pdfmetrics.getRegisteredFontNames():
        return font_name
    font_path = _find_font_path()
    if font_path:
        pdfmetrics.registerFont(TTFont(font_name, str(font_path)))
        return font_name
    return "Helvetica"


def _pdf_styles() -> dict[str, ParagraphStyle]:
    font_name = _register_pdf_font()
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("title_cn", parent=base["Title"], fontName=font_name, fontSize=22, leading=28, alignment=TA_CENTER, textColor=colors.HexColor("#0B2545"), spaceAfter=12),
        "subtitle": ParagraphStyle("subtitle_cn", parent=base["Normal"], fontName=font_name, fontSize=12, leading=17, alignment=TA_CENTER, textColor=colors.HexColor("#2E74B5"), spaceAfter=16),
        "h1": ParagraphStyle("h1_cn", parent=base["Heading1"], fontName=font_name, fontSize=14, leading=19, textColor=colors.HexColor("#2E74B5"), spaceBefore=10, spaceAfter=8),
        "body": ParagraphStyle("body_cn", parent=base["Normal"], fontName=font_name, fontSize=9.5, leading=15, alignment=TA_LEFT, spaceAfter=5),
        "small": ParagraphStyle("small_cn", parent=base["Normal"], fontName=font_name, fontSize=8, leading=12, alignment=TA_LEFT),
    }


def _p(text: Any, style: ParagraphStyle) -> Paragraph:
    return Paragraph(html.escape(_safe_text(text)).replace("\n", "<br/>"), style)


def _pdf_cell_text(value: Any, max_chars: int = 520) -> str:
    text = _safe_text(value).strip()
    if len(text) > max_chars:
        return text[:max_chars].rstrip() + "……"
    return text


def _pdf_table(headers: list[str], rows: list[list[Any]], styles: dict[str, ParagraphStyle], max_rows: int = 80) -> Table:
    data = [[_p(header, styles["small"]) for header in headers]]
    for row in rows[:max_rows]:
        data.append([_p(_pdf_cell_text(value), styles["small"]) for value in row])
    table = LongTable(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F0F6F9")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#0B2545")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D9E5EE")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    return table


def build_pdf_report(data: ReportData) -> bytes:
    payload = data.payload
    chart_payloads = _chart_payloads(payload)
    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=14 * mm, bottomMargin=14 * mm)
    styles = _pdf_styles()
    story: list[Any] = [
        Paragraph("数据智能体分析报告", styles["title"]),
        Paragraph(_report_title(data), styles["subtitle"]),
        Paragraph("一、报告概况", styles["h1"]),
        _pdf_table(
            ["项目", "内容"],
            [
                ["会话编号", data.session["id"]],
                ["分析问题", data.question],
                ["分析类型", payload.get("intent", "")],
                ["执行模式", payload.get("execution_mode", "")],
                ["生成时间", data.assistant_message.get("created_at", "")],
            ],
            styles,
        ),
        Spacer(1, 8),
        Paragraph("二、回答与关键发现", styles["h1"]),
    ]
    for idx, insight in enumerate(_top_report_insights(data, chart_payloads), 1):
        story.append(Paragraph(f"{idx}. {html.escape(_safe_text(insight))}", styles["body"]))

    temp_chart_paths: list[str] = []
    if chart_payloads:
        story.extend([Spacer(1, 8), Paragraph("三、结果图表", styles["h1"])])
        for index, chart_payload in enumerate(chart_payloads, 1):
            chart_image = build_chart_image(chart_payload)
            if not chart_image:
                continue
            story.append(Paragraph(f"{index}. {html.escape(_safe_text(chart_payload.get('title') or '图表'))}", styles["body"]))
            tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
            tmp.write(chart_image.getvalue())
            tmp.close()
            temp_chart_paths.append(tmp.name)
            story.append(PdfImage(tmp.name, width=245 * mm, height=116 * mm))
            chart_insights = _chart_insights(chart_payload, payload)
            if chart_insights:
                story.append(Paragraph("该图提示：", styles["body"]))
                for insight_index, insight in enumerate(chart_insights, 1):
                    story.append(Paragraph(f"{insight_index}. {html.escape(_safe_text(insight))}", styles["small"]))
            story.append(Spacer(1, 6))

    next_section = 4 if chart_payloads else 3
    overall = _overall_insights(payload, chart_payloads)
    if overall:
        story.extend([Spacer(1, 8), Paragraph(_section_heading(next_section, "总体建议"), styles["h1"])])
        next_section += 1
        for idx, insight in enumerate(overall, 1):
            story.append(Paragraph(f"{idx}. {html.escape(_safe_text(insight))}", styles["body"]))

    if payload.get("rows"):
        story.extend([PageBreak(), Paragraph(_section_heading(next_section, "查询结果"), styles["h1"])])
        next_section += 1
        columns = payload.get("columns", [])
        all_rows = payload.get("rows", [])
        max_show = 10
        story.append(_pdf_table(columns, [[row.get(column, "") for column in columns] for row in all_rows[:max_show]], styles, max_rows=max_show))
        if len(all_rows) > max_show:
            story.append(Paragraph(f"（结果共 {len(all_rows)} 行，PDF 报告仅展示最重要的前 {max_show} 行。）", styles["small"]))
    if payload.get("sql"):
        story.extend([Spacer(1, 8), Paragraph(_section_heading(next_section, "执行 SQL"), styles["h1"]), _p(payload["sql"], styles["small"])])
        next_section += 1
    if payload.get("knowledge_refs"):
        story.extend([Spacer(1, 8), Paragraph(_section_heading(next_section, "知识依据"), styles["h1"])])
        story.append(
            _pdf_table(
                ["标题", "类别"],
                [[_compact_reference_title(item.get("title", "")), item.get("category", "")] for item in payload["knowledge_refs"]],
                styles,
                max_rows=20,
            )
        )
    try:
        doc.build(story)
    finally:
        for temp_chart_path in temp_chart_paths:
            try:
                os.remove(temp_chart_path)
            except OSError:
                pass
    return out.getvalue()


def build_multi_section_pdf(
    report_title: str,
    sections: list[dict],
    executive_summary: str = "",
    data_source: str = "",
    sql_list: list[str] | None = None,
) -> bytes:
    """Phase 3: Multi-section report PDF with inline charts per section."""
    out = BytesIO()
    doc = SimpleDocTemplate(out, pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=18 * mm, bottomMargin=18 * mm)
    styles = _pdf_styles()
    story: list[Any] = []

    # Cover
    story.append(Spacer(1, 40 * mm))
    story.append(Paragraph("数据智能体分析报告", styles["title"]))
    story.append(Spacer(1, 8 * mm))
    story.append(Paragraph(report_title[:120], styles["subtitle"]))
    story.append(Spacer(1, 12 * mm))
    if data_source:
        story.append(Paragraph(f"数据来源: {data_source}", styles["body"]))
    from datetime import datetime
    story.append(Paragraph(f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles["body"]))
    story.append(PageBreak())

    # Executive Summary
    if executive_summary:
        story.append(Paragraph("执行摘要", styles["h1"]))
        for line in executive_summary.split("\n"):
            line = line.strip()
            if line:
                story.append(Paragraph(line, styles["body"]))
        story.append(PageBreak())

    # Sections with inline charts
    nums = "一二三四五六七八九十"
    for idx, section in enumerate(sections, 1):
        title = section.get("title", f"分析维度 {idx}")
        narrative = section.get("narrative", "")
        rows = section.get("rows", [])
        chart_bytes = section.get("chart_image")

        cn = nums[idx - 1] if 1 <= idx <= 10 else str(idx)
        story.append(Paragraph(f"{cn}、{title}", styles["h1"]))

        if chart_bytes:
            try:
                tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
                data = chart_bytes if isinstance(chart_bytes, bytes) else (chart_bytes.getvalue() if hasattr(chart_bytes, 'getvalue') else b'')
                tmp.write(data)
                tmp.close()
                story.append(PdfImage(tmp.name, width=160 * mm, height=75 * mm))
                story.append(Spacer(1, 4 * mm))
                os.remove(tmp.name)
            except Exception:
                pass

        for line in narrative.split("\n"):
            line = line.strip()
            if line:
                story.append(Paragraph(line, styles["body"]))
        story.append(Spacer(1, 4 * mm))

        if rows and len(rows) <= 10:
            cols = list(rows[0].keys())[:6]
            story.append(_pdf_table(cols, [[str(r.get(c, "") or "")[:40] for c in cols] for r in rows[:10]], styles, max_rows=10))
        elif rows:
            cols = list(rows[0].keys())[:6]
            story.append(_pdf_table(cols, [[str(r.get(c, "") or "")[:40] for c in cols] for r in rows[:10]], styles, max_rows=10))
            story.append(Paragraph(f"（该分析维度共 {len(rows)} 条记录，PDF 报告仅展示最重要的前 10 行。）", styles["small"]))
        story.append(PageBreak())

    # Appendix
    if sql_list:
        story.append(Paragraph("附录: 分析 SQL", styles["h1"]))
        for i, sql in enumerate(sql_list[:10], 1):
            story.append(Paragraph(f"{i}. {sql[:200]}", styles["small"]))
            story.append(Spacer(1, 2 * mm))

    doc.build(story)
    return out.getvalue()
